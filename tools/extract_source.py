#!/usr/bin/env python3
"""Extract one source document into working text + images.
Usage: python3 tools/extract_source.py <session LNN> <source file> [<source file> ...]
Output: Resources/extract/LNN/<slug>/text.md and images/ (gitignored — raw material, never published).
Owner skill: dgs-intake. Condensing into knowledge_base/ is a separate, human-reviewed step."""
import hashlib, html, re, subprocess, sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Resources" / "extract"

def slug(name):
    s = re.sub(r"[^A-Za-z0-9]+", "-", name).strip("-").lower()
    return s[:60] or "doc"

def extract_pptx(src, dest):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(str(src))
    imgdir = dest / "images"; imgdir.mkdir(parents=True, exist_ok=True)
    seen = {}; lines = [f"# {src.name}", f"_source: {src}_  \n_slides: {len(prs.slides)}_\n"]
    n_img = 0
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        lines.append(f"\n## Slide {i}" + (f" — {title}" if title else ""))
        for sh in slide.shapes:
            if sh.has_text_frame and sh != slide.shapes.title:
                for p in sh.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        lines.append(("  " * min(p.level, 3)) + "- " + t)
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    lines.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    blob = sh.image.blob; h = hashlib.sha1(blob).hexdigest()[:10]
                    ext = sh.image.ext
                    if h not in seen:
                        fn = f"s{i:03d}_{h}.{ext}"; (imgdir / fn).write_bytes(blob); seen[h] = fn; n_img += 1
                    lines.append(f"![img](images/{seen[h]})")
                except Exception:
                    pass
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("> NOTES: " + notes.replace("\n", "\n> "))
    (dest / "text.md").write_text("\n".join(lines), encoding="utf-8")
    return len(prs.slides), n_img

def extract_pdf(src, dest):
    imgdir = dest / "images"; imgdir.mkdir(parents=True, exist_ok=True)
    txt = subprocess.run(["pdftotext", "-layout", str(src), "-"], capture_output=True, text=True).stdout
    pages = txt.split("\f")
    lines = [f"# {src.name}", f"_source: {src}_  \n_pages: {len(pages)}_\n"]
    for i, pg in enumerate(pages, 1):
        pg = pg.strip()
        if pg:
            lines.append(f"\n## Page {i}\n\n```\n{pg}\n```")
    (dest / "text.md").write_text("\n".join(lines), encoding="utf-8")
    subprocess.run(["pdfimages", "-png", "-p", str(src), str(imgdir / "p")], capture_output=True)
    imgs = list(imgdir.glob("*.png"))
    # drop tiny decorations
    for p in imgs:
        if p.stat().st_size < 4000:
            p.unlink()
    return len(pages), len(list(imgdir.glob("*.png")))

def extract_html(src, dest):
    dest.mkdir(parents=True, exist_ok=True)
    raw = src.read_text(encoding="utf-8", errors="ignore")
    (dest / "source.html").write_text(raw, encoding="utf-8")       # keep it: the diagrams live here
    svgs = re.findall(r"<svg.*?</svg>", raw, flags=re.S | re.I)
    for k, s in enumerate(svgs, 1):
        (dest / f"diagram_{k:02d}.svg").write_text(s, encoding="utf-8")
    body = re.sub(r"<(script|style).*?</\1>", "", raw, flags=re.S | re.I)
    body = re.sub(r"<h([1-4])[^>]*>(.*?)</h\1>", lambda m: "\n" + "#" * int(m.group(1)) + " " + m.group(2) + "\n", body, flags=re.S | re.I)
    body = re.sub(r"<li[^>]*>", "\n- ", body, flags=re.I)
    body = re.sub(r"<(p|div|br|tr)[^>]*>", "\n", body, flags=re.I)
    text = html.unescape(re.sub(r"<[^>]+>", "", body))
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    (dest / "text.md").write_text(f"# {src.name}\n_source: {src}_\n\n" + text.strip(), encoding="utf-8")
    return text.count("\n"), len(svgs)

def main():
    if len(sys.argv) < 3:
        print(__doc__); sys.exit(1)
    sess = sys.argv[1].upper()
    for arg in sys.argv[2:]:
        src = Path(arg)
        if not src.exists():
            print(f"  MISSING {src}"); continue
        dest = OUT / sess / slug(src.stem); dest.mkdir(parents=True, exist_ok=True)
        ext = src.suffix.lower()
        if ext == ".pptx":   n, k = extract_pptx(src, dest); kind = "slides"
        elif ext == ".pdf":  n, k = extract_pdf(src, dest);  kind = "pages"
        elif ext in (".html", ".htm"): n, k = extract_html(src, dest); kind = "lines"
        else:
            print(f"  SKIP {src.name} (unsupported)"); continue
        print(f"  {sess}  {src.name:60s} {n:4d} {kind:6s} {k:3d} images/svg -> {dest.relative_to(ROOT)}")

if __name__ == "__main__":
    main()
